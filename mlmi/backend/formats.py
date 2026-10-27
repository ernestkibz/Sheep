"""Sheep's own document formats, and the only doors to Microsoft's.

The AI and the editors work on these; .docx/.xlsx/.pptx come IN through
import and go OUT through export, nowhere else (Ernest, 2026-09-30). Each is
one JSON file a model can read and a herder can transform:

  .sdoc    Sheep Words          {"sheep": "words", "blocks": [...]}
  .ssheet  Sheep Sheets         {"sheep": "sheets", "sheets": [{"name", "cells", "fmt", ...}]}
  .sdeck   Sheep Presentations  {"sheep": "presentations", "theme", "slides": [...]}

Text carries its formatting inline -- **bold**, *italic*, __underline__ -- so a
paragraph is one string a small model can write, not a list of runs.
Pictures are workspace files, referenced by path; importing a document
extracts its embedded pictures into a folder beside it.
"""
import json
import os
import re
import time

EXT = {".sdoc": "words", ".ssheet": "sheets", ".sdeck": "presentations", ".sstudio": "studio"}
_MARKER = re.compile(r"^\s*(?:[\u2022\u00b7\u25aa\u25cf\u2013\u2014*-]|\d{1,2}[.)])\s+")


def clean(text):
    """Text as a small model sends it, made plain: escape sequences it wrote out
    literally (a backslash and an n, t or u2022) and tab runs become what they mean."""
    s = str(text or "")
    s = s.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\\t", " ").replace("\\u2022", "\u2022")
    s = s.replace("\r\n", "\n").replace("\t", " ")
    return "\n".join(line.strip() for line in s.split("\n")).strip()


def items(text):
    """One list item per line: a model that put the whole list in one call gets a list."""
    return [item(x) for x in clean(text).split("\n") if item(x).strip()]


def item(text):
    """A list item without the bullet or number a model typed in front of it
    ("\u2022 milk", "- milk", "1. milk"): the block draws its own marker. Seen on
    the first Llama 3.2 3B run, which wrote "\u2022 milk" into a bullet block."""
    return _MARKER.sub("", str(text or ""), count=1)
OFFICE = {".docx": ".sdoc", ".xlsx": ".ssheet", ".pptx": ".sdeck"}
THEMES = {
    # name: (background, title colour, body colour, accent)
    "light":    ("FFFFFF", "1F2937", "374151", "2563EB"),
    "dark":     ("111827", "F9FAFB", "D1D5DB", "60A5FA"),
    "midnight": ("0B1020", "E0E7FF", "C7D2FE", "A78BFA"),
    "forest":   ("F1F5EE", "1E3A2B", "2F4A3A", "3F7D4E"),
    "sunset":   ("FFF7ED", "7C2D12", "9A3412", "EA580C"),
    "mono":     ("FAFAFA", "0A0A0A", "262626", "737373"),
}


def new(kind, **kw):
    if kind == "words":
        return {"sheep": "words", "v": 1, "blocks": []}
    if kind == "sheets":
        return {"sheep": "sheets", "v": 1, "active": 0,
                "sheets": [{"name": kw.get("sheet") or "Sheet1", "cells": {}, "fmt": {}, "widths": {}, "charts": []}]}
    if kind == "presentations":
        return {"sheep": "presentations", "v": 1, "theme": kw.get("theme") or "light", "ratio": [16, 9], "slides": []}
    raise ValueError(kind)


def load(p):
    with open(p, encoding="utf-8") as f:
        d = json.load(f)
    if not isinstance(d, dict) or d.get("sheep") not in ("words", "sheets", "presentations", "studio"):
        raise ValueError("not a Sheep document")
    return d


def save(p, d):
    tmp = str(p) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=1)
    # Windows refuses the replace while another reader holds the file open: the
    # editor re-reads a document on every change event, and a scanner may peek.
    # A Words job stopped on "[WinError 5] Access is denied" mid-build in the
    # 2026-10-12 simulation. Retry the replace briefly, as save_json does.
    for attempt in range(10):
        try:
            os.replace(tmp, p)
            return
        except PermissionError:
            if attempt == 9:
                try:
                    os.remove(tmp)
                except OSError:
                    pass
                raise
            time.sleep(0.02 * (attempt + 1))


# ------------------------------------------------------------ inline marks

MARK = re.compile(r"(\\[*_\\]|\*\*|__|\*)")


def runs(text):
    """'a **b** *c* __d__' -> [(text, bold, italic, underline)]"""
    out, b, i, u, buf = [], False, False, False, ""
    for part in MARK.split(str(text or "")):
        if part in ("**", "*", "__"):
            if buf:
                out.append((buf, b, i, u))
                buf = ""
            if part == "**":
                b = not b
            elif part == "*":
                i = not i
            else:
                u = not u
        elif part.startswith("\\") and len(part) == 2:
            buf += part[1]
        else:
            buf += part
    if buf:
        out.append((buf, b, i, u))
    return out


def marks(rs):
    """[(text, bold, italic, underline)] -> inline-marked text"""
    s = ""
    for t, b, i, u in rs:
        t = t.replace("\\", "\\\\").replace("*", "\\*").replace("__", "\\_\\_")
        core = t.strip()
        if not core or not (b or i or u):
            s += t
            continue
        lead, trail = t[:len(t) - len(t.lstrip())], t[len(t.rstrip()):]
        o = ("**" if b else "") + ("*" if i else "") + ("__" if u else "")
        c = ("__" if u else "") + ("*" if i else "") + ("**" if b else "")
        s += lead + o + core + c + trail
    return s


def plain(text):
    return "".join(r[0] for r in runs(text))


# ================================================================== Words

def docx_import(src, assets_dir, rel_assets):
    """.docx -> .sdoc blocks. Pictures are written into assets_dir."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(src)
    blocks, n_img = [], 0
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:tbl"):
            t = Table(el, doc._body)
            blocks.append({"type": "table", "header": True,
                           "rows": [[marks([(c.text, False, False, False)]) for c in r.cells] for r in t.rows]})
            continue
        if el.tag != qn("w:p"):
            continue
        p = Paragraph(el, doc._body)
        # pictures inside the paragraph become their own blocks
        for blip in el.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            part = doc.part.related_parts.get(rid)
            if part is None:
                continue
            n_img += 1
            ext = os.path.splitext(part.partname)[1] or ".png"
            os.makedirs(assets_dir, exist_ok=True)
            name = f"image{n_img}{ext}"
            with open(os.path.join(assets_dir, name), "wb") as f:
                f.write(part.blob)
            blocks.append({"type": "image", "src": f"{rel_assets}/{name}", "width": 5.0})
        style = (p.style.name if p.style is not None else "Normal") or "Normal"
        text = marks([(r.text, bool(r.bold), bool(r.italic), bool(r.underline)) for r in p.runs])
        if not text.strip() and el.findall(".//" + qn("a:blip")):
            continue
        if any(br.get(qn("w:type")) == "page" for br in el.findall(".//" + qn("w:br"))) and not text.strip():
            blocks.append({"type": "pagebreak"})
            continue
        b = {"text": text}
        lvl = el.find(".//" + qn("w:ilvl"))
        if style == "Title":
            b.update(type="heading", level=0)
        elif style.startswith("Heading"):
            try:
                b.update(type="heading", level=int(style.split()[-1]))
            except ValueError:
                b.update(type="heading", level=1)
        elif "Number" in style:
            b.update(type="numbered", level=int(lvl.get(qn("w:val"))) if lvl is not None else 0)
        elif "Bullet" in style or "List" in style or el.find(".//" + qn("w:numPr")) is not None:
            b.update(type="bullet", level=int(lvl.get(qn("w:val"))) if lvl is not None else 0)
        else:
            b.update(type="paragraph", style={"Quote": "quote", "Subtitle": "subtitle",
                                              "Caption": "caption"}.get(style, "normal"))
        if p.alignment is not None:
            a = str(p.alignment).split(".")[-1].split(" ")[0].lower()
            if a in ("center", "right", "justify"):
                b["align"] = a
        blocks.append(b)
    title = next((plain(b["text"]) for b in blocks if b.get("type") == "heading" and b.get("level") == 0), "")
    return {"sheep": "words", "v": 1, "title": title, "blocks": blocks}


def docx_export(d, dst, ws_root):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A, WD_BREAK
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    def fill(par, text, blk):
        for t, b, i, u in runs(text):
            r = par.add_run(t)
            r.bold, r.italic, r.underline = b or None, i or None, u or None
            if blk.get("size"):
                r.font.size = Pt(float(blk["size"]))
            if blk.get("color"):
                try:
                    r.font.color.rgb = RGBColor.from_string(str(blk["color"]).strip("#")[:6].upper())
                except ValueError:
                    pass
        if blk.get("align"):
            par.alignment = {"left": A.LEFT, "center": A.CENTER, "right": A.RIGHT, "justify": A.JUSTIFY}.get(blk["align"])

    for b in d.get("blocks", []):
        t = b.get("type")
        if t == "heading":
            h = doc.add_heading("", max(0, min(int(b.get("level", 1)), 9)))
            fill(h, b.get("text", ""), b)
        elif t in ("bullet", "numbered"):
            base = "List Bullet" if t == "bullet" else "List Number"
            lv = int(b.get("level", 0))
            style = base if lv == 0 else f"{base} {min(lv + 1, 3)}"
            fill(doc.add_paragraph(style=style), b.get("text", ""), b)
        elif t == "table":
            rows = b.get("rows") or [[""]]
            cols = max(len(r) for r in rows)
            tb = doc.add_table(rows=len(rows), cols=cols)
            tb.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    par = tb.cell(ri, ci).paragraphs[0]
                    text = row[ci] if ci < len(row) else ""
                    if ri == 0 and b.get("header", True):
                        text = f"**{plain(text)}**" if plain(text) else ""
                    fill(par, text, {})
        elif t == "image":
            src = os.path.join(ws_root, b.get("src", ""))
            if os.path.exists(src):
                doc.add_picture(src, width=Inches(float(b.get("width") or 5.0)))
        elif t == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            style = {"quote": "Quote", "subtitle": "Subtitle", "caption": "Caption"}.get(b.get("style"), "Normal")
            fill(doc.add_paragraph(style=style), b.get("text", ""), b)
    doc.save(dst)


def words_outline(d, limit=150):
    lines = []
    for i, b in enumerate(d.get("blocks", [])):
        t = b.get("type")
        if t == "table":
            rows = b.get("rows") or []
            lines.append(f"[{i}] table {len(rows)}x{max((len(r) for r in rows), default=0)}: "
                         + " | ".join(plain(c)[:18] for c in (rows[0] if rows else [])))
        elif t == "image":
            lines.append(f"[{i}] image: {b.get('src')}")
        elif t == "pagebreak":
            lines.append(f"[{i}] page break")
        else:
            lab = ("title" if b.get("level") == 0 else f"heading {b.get('level', 1)}") if t == "heading" else \
                (f"{t}" + (f" (level {b['level']})" if b.get("level") else "") if t in ("bullet", "numbered")
                 else b.get("style", "paragraph") if b.get("style", "normal") != "normal" else "paragraph")
            txt = b.get("text", "")
            lines.append(f"[{i}] {lab}: {txt[:limit]}{'...' if len(txt) > limit else ''}")
    return "\n".join(lines) if lines else "(empty document)"


# ================================================================== Sheets

def xlsx_import(src):
    from openpyxl import load_workbook
    wb = load_workbook(src)
    out = {"sheep": "sheets", "v": 1, "active": wb.sheetnames.index(wb.active.title), "sheets": []}
    for ws in wb.worksheets:
        cells, fmt = {}, {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    v = c.value
                    if not isinstance(v, (int, float, str, bool)):
                        v = str(v)                       # dates and times come in as their text
                    cells[c.coordinate] = v
                f = {}
                if c.font is not None and c.font.b:
                    f["b"] = True
                if c.fill is not None and c.fill.fill_type == "solid" and c.fill.fgColor.rgb not in (None, "00000000") \
                        and isinstance(c.fill.fgColor.rgb, str):
                    f["bg"] = c.fill.fgColor.rgb[-6:]
                if c.number_format and c.number_format != "General":
                    f["nf"] = c.number_format
                if c.alignment is not None and c.alignment.horizontal in ("left", "center", "right"):
                    f["align"] = c.alignment.horizontal
                if f:
                    fmt[c.coordinate] = f
        widths = {k: v.width for k, v in ws.column_dimensions.items() if v.width and v.customWidth}
        out["sheets"].append({"name": ws.title, "cells": cells, "fmt": fmt, "widths": widths, "charts": []})
    return out


def xlsx_export(d, dst):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import range_boundaries
    wb = Workbook()
    wb.remove(wb.active)
    for sh in d.get("sheets", []):
        ws = wb.create_sheet(sh.get("name", "Sheet")[:31])
        for k, v in sh.get("cells", {}).items():
            ws[k] = v
        for k, f in sh.get("fmt", {}).items():
            c = ws[k]
            if f.get("b"):
                c.font = Font(bold=True)
            if f.get("bg"):
                c.fill = PatternFill("solid", fgColor=f["bg"])
            if f.get("nf"):
                c.number_format = f["nf"]
            if f.get("align"):
                c.alignment = Alignment(horizontal=f["align"])
        for col, w in sh.get("widths", {}).items():
            ws.column_dimensions[col].width = float(w)
        for ch in sh.get("charts", []):
            try:
                c0, r0, c1, r1 = range_boundaries(ch["range"])
            except (KeyError, ValueError, TypeError):
                continue
            chart = {"line": LineChart, "pie": PieChart}.get(ch.get("kind"), BarChart)()
            chart.title = ch.get("title") or None
            chart.add_data(Reference(ws, min_col=c0 + 1, min_row=r0, max_col=c1, max_row=r1), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=c0, min_row=r0 + 1, max_row=r1))
            chart.width, chart.height = 16, 8
            ws.add_chart(chart, ch.get("anchor") or "H2")
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.active = min(int(d.get("active", 0)), len(wb.sheetnames) - 1)
    wb.save(dst)


class SheetView:
    """What the formula evaluator reads: cells by sheet name, over a .ssheet."""

    def __init__(self, d):
        self.by = {s["name"]: s for s in d.get("sheets", [])}
        self.sheetnames = list(self.by)

    def get(self, sheet, coord):
        return self.by[sheet]["cells"].get(coord)


# ================================================================== Presentations

def pptx_import(src, assets_dir, rel_assets):
    from pptx import Presentation
    prs = Presentation(src)
    W, H = prs.slide_width, prs.slide_height
    theme = ""
    cat = prs.core_properties.category or ""
    if cat.startswith("sheep-theme:"):
        theme = cat.split(":", 1)[1]
    out = {"sheep": "presentations", "v": 1, "theme": theme or "light",
           "ratio": [16, 9] if abs(W / H - 16 / 9) < 0.05 else [4, 3], "slides": []}
    n_img = 0
    for s in prs.slides:
        sl = {"layout": "bullets", "title": "", "bullets": [], "notes": "", "items": []}
        for sh in s.shapes:
            if sh.is_placeholder and sh.has_text_frame:
                idx = sh.placeholder_format.idx
                typ = str(sh.placeholder_format.type).split(" ")[0]
                if typ == "SUBTITLE":            # before the title test: "SUBTITLE" contains "TITLE"
                    sl["subtitle"] = sh.text_frame.text
                    continue
                if idx == 0 or typ in ("TITLE", "CENTER_TITLE"):
                    sl["title"] = sh.text_frame.text
                    if typ == "CENTER_TITLE":
                        sl["layout"] = "title"
                    continue
                if idx == 1:
                    sl["bullets"] = ["  " * p.level + p.text for p in sh.text_frame.paragraphs if p.text.strip()]
                    continue
            if sh.left is None:
                continue
            box = {"x": round(sh.left / W * 100, 2), "y": round(sh.top / H * 100, 2),
                   "w": round(sh.width / W * 100, 2), "h": round(sh.height / H * 100, 2)}
            if sh.shape_type == 13:
                n_img += 1
                os.makedirs(assets_dir, exist_ok=True)
                name = f"slide{len(out['slides']) + 1}_{n_img}.{sh.image.ext}"
                with open(os.path.join(assets_dir, name), "wb") as f:
                    f.write(sh.image.blob)
                sl["items"].append({"kind": "image", "src": f"{rel_assets}/{name}", **box})
            elif sh.has_text_frame and sh.text_frame.text.strip():
                size = next((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), None)
                sl["items"].append({"kind": "text", "text": sh.text_frame.text, "size": size, **box})
        if s.has_notes_slide:
            sl["notes"] = s.notes_slide.notes_text_frame.text
        if not sl["bullets"] and sl["layout"] != "title" and not sl["items"]:
            sl["layout"] = "title_only"
        out["slides"].append(sl)
    return out


def pptx_export(d, dst, ws_root):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    prs = Presentation()
    wide = list(d.get("ratio", [16, 9])) == [16, 9]
    if wide:
        # the built-in template is 4:3: widen its layouts with the slide
        k = 13.333 / 10.0
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for shp in list(prs.slide_master.shapes) + [s for l in prs.slide_layouts for s in l.placeholders]:
            if shp.left is not None and shp.width is not None:
                shp.left, shp.width = int(shp.left * k), int(shp.width * k)
    theme = d.get("theme") if d.get("theme") in THEMES else "light"
    prs.core_properties.category = "sheep-theme:" + theme
    bg, tc, bc, _ = THEMES[theme]
    W, H = prs.slide_width, prs.slide_height
    lay = {"title": 0, "bullets": 1, "section": 2, "title_only": 5, "blank": 6}
    for sl in d.get("slides", []):
        s = prs.slides.add_slide(prs.slide_layouts[lay.get(sl.get("layout"), 1)])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
        if s.shapes.title is not None:
            s.shapes.title.text = plain(sl.get("title", ""))
        body = next((p for p in s.placeholders if p.placeholder_format.idx == 1), None)
        bullets = sl.get("bullets") or []
        if body is not None:
            if sl.get("layout") == "title":
                body.text = plain(sl.get("subtitle", ""))
            elif bullets:
                tf = body.text_frame
                tf.clear()
                for k2, bt in enumerate(bullets):
                    lvl = (len(bt) - len(bt.lstrip(" "))) // 2
                    par = tf.paragraphs[0] if k2 == 0 else tf.add_paragraph()
                    par.text = plain(bt.strip().lstrip("-").strip())
                    par.level = min(lvl, 4)
                if any(i.get("kind") == "image" for i in sl.get("items", [])):
                    body.width = int(W * 0.5)
            else:
                body._element.getparent().remove(body._element)
        for it in sl.get("items", []):
            x, y = int(W * float(it.get("x", 10)) / 100), int(H * float(it.get("y", 20)) / 100)
            w = int(W * float(it.get("w", 40)) / 100)
            if it.get("kind") == "image":
                src = os.path.join(ws_root, it.get("src", ""))
                if os.path.exists(src):
                    s.shapes.add_picture(src, x, y, width=w)
            elif it.get("kind") == "text":
                h = int(H * float(it.get("h", 10)) / 100)
                tb = s.shapes.add_textbox(x, y, w, h)
                tb.text_frame.word_wrap = True
                tb.text_frame.text = plain(it.get("text", ""))
                for par in tb.text_frame.paragraphs:
                    for r in par.runs:
                        if it.get("size"):
                            r.font.size = Pt(float(it["size"]))
                        if it.get("b"):
                            r.font.bold = True
        for sh in s.shapes:                               # the theme's text colours
            if sh.has_text_frame:
                is_title = sh == s.shapes.title
                for par in sh.text_frame.paragraphs:
                    for r in par.runs:
                        r.font.color.rgb = RGBColor.from_string(tc if is_title else bc)
                        if is_title and sl.get("title_scale"):      # the theme's 40/44 pt, scaled
                            base = 44 if sl.get("layout") in ("title", "section") else 40
                            r.font.size = Pt(round(base * float(sl["title_scale"])))
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]
    prs.save(dst)


def deck_outline(d):
    lines = [f"{len(d.get('slides', []))} slides, theme {d.get('theme', 'light')}"]
    for i, s in enumerate(d.get("slides", []), 1):
        extra = [it for it in s.get("items", [])]
        lines.append(f"slide {i}: {s.get('title') or '(no title)'}" + (f"  [{s.get('layout')}]" if s.get("layout") not in (None, "bullets") else "")
                     + (f"  [{len(extra)} item(s)]" if extra else ""))
        if s.get("subtitle"):
            lines.append(f"   subtitle: {s['subtitle']}")
        lines.extend("   " + ("  " * ((len(b) - len(b.lstrip(' '))) // 2)) + "- " + b.strip()[:100]
                     for b in (s.get("bullets") or [])[:12])
        if s.get("notes"):
            lines.append(f"   notes: {s['notes'][:120]}")
    return "\n".join(lines)
