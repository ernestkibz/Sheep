"""What a turn may change: in-scope checks, artifact guards, the honesty phrases, plain errors.

Split out of agent.py (2026-10-21); every name is re-exported from the agent package."""
import re
from .. import tools as T
from .. import checker, cues, shepherd, verify
from ..workspace import ModelSelectionRequired, ToolError


def _content(name, args):
    """What an adding call would put in a file, as (path, text) pairs; empty for
    other calls. The 3B read its new list back and added the same three bullets
    again (2026-09-30): a job may not add what it already added."""
    from .. import formats as F
    t = T.TOOLS.get(name)
    if not t:
        return frozenset()
    path = str(args.get("path") or "").lower()
    if name == "words_add":
        texts = F.items(args.get("text")) or [F.clean(args.get("text"))]
    elif name == "words_add_list":
        items = T._coerce(t.params["items"], args.get("items"))
        texts = [x for it in (items if isinstance(items, list) else [items]) for x in F.items(it)]
    elif name == "slides_add":
        texts = ["slide: " + F.clean(args.get("title"))]
    else:
        return frozenset()
    return frozenset((path, x.lower()) for x in texts if x)


def _overwrites_own(ws, name, args, made):
    """A sheet write that would change cells this same job already filled. Seen
    on Sheep + Nemotron (2026-09-30): the plan added writes after sheets_create
    had made the whole table, and "Total, 155" landed on the Fuel row. A job's
    own new workbook is complete once written; changing it again is not asked."""
    if name not in ("sheets_write", "sheets_set_cell"):
        return None
    path = str(args.get("path") or "").replace("\\", "/")
    if path not in made:
        return None
    try:
        from openpyxl.utils import get_column_letter
        from .. import formats as F
        from ..tools import sheets as SH
        d = F.load(ws.path(path))
        sheet = SH._sheet(d, args.get("sheet") or "")
        c0, r0 = SH._cell(args.get("cell"))
        vals = [[args.get("value")]] if name == "sheets_set_cell" else T._coerce(("array", ""), args.get("values"))
        rows = [r if isinstance(r, list) else [r] for r in (vals if isinstance(vals, list) else [[vals]])]
    except Exception:
        return None
    for dr, row in enumerate(rows):
        for dc, v in enumerate(row):
            k = f"{get_column_letter(c0 + dc)}{r0 + dr}"
            cur = sheet["cells"].get(k)
            if cur not in (None, "") and str(SH.coerce(v)) != str(cur):
                return (f"{k} already holds {cur!r}, written earlier in this job; the table is complete. "
                        "The job is done when the request is done.")
    return None


class OutOfScope(ToolError):
    """A write the request did not ask for: the job ends here."""


def _norm(x):
    import re
    return re.sub(r"[\s_.\-/\\]+", " ", str(x or "").lower()).strip()


def in_scope(ws, name, args, asked, open_path, made):
    """The scope guard (Ernest: do for the user what they asked, nothing else).
    Llama 3.2 3B, after making the list it was asked for, went on to create
    budget.ssheet.sdoc and report.q3.sdoc and to write into the person's own
    Project brief (2026-09-30). A write may touch only: the file open in the
    editor, a file the request names, a file this job made, or -- once per job
    -- one new file (the thing the person asked to be made). Reads are free."""
    t = T.TOOLS.get(name)
    if not t or t.tier < 2 or name.startswith("herder_") and name in ("herder_help", "herder_list", "herder_show"):
        return None
    targets = [args.get(k) for k in ("path", "src", "dst") if args.get(k)]
    a = _norm(asked)
    for tg in targets:
        rel = str(tg).replace("\\", "/")
        if rel in made or (open_path and _norm(rel) == _norm(open_path)):
            continue
        stem = rel.rsplit("/", 1)[-1].rsplit(".", 1)[0]
        if _norm(rel) in a or (len(_norm(stem)) >= 3 and _norm(stem) in a):
            continue
        try:
            exists = ws.path(rel).exists()
        except ToolError:
            exists = False
        if not exists and not made:
            continue
        return (f"{rel} is not part of this request" + ("" if exists else ": this job already made its file")
                + ". The job is done when the request is done.")
    return None


def _header(ws, path):
    """A sheet's first row, as text, in column order."""
    from .. import formats as F
    try:
        cells = F.load(ws.path(path, must_exist=True))["sheets"][0]["cells"]
    except Exception:
        return []
    row = sorted(((k, v) for k, v in cells.items() if re.fullmatch(r"[A-Z]+1", k)), key=lambda kv: (len(kv[0]), kv[0]))
    return [str(v.get("v") if isinstance(v, dict) else v) for _, v in row]


# a question about a file, and the words that make a request change one
READ_ONLY = (r"\b(summari[sz]e|summary of|what (does|is|are|did|was|were|do)|how (many|much|long)|tell me|explain|"
             r"describe|read|which|who|when|where|count|total|average|sum of|check|look at|review)\b")


WRITE_WORDS = (r"\b(add|create|make|write|insert|change|update|edit|delete|remove|rename|move|copy|export|replace|"
               r"fix|set|sort|put|save|format|append|fill|convert|turn)\b")


# "what did I save here?" is a question, not an instruction to save, so the veto
# for this one path leaves "save" and "store" out.
LIST_WRITE = (r"\b(add|create|make|write|insert|change|update|edit|delete|remove|rename|move|copy|export|"
              r"import|replace|fix|sort|format|append|fill|convert)\b")


# a reply that says it did file work ("I added a new slide", "Here's the updated file")
# a reply that says the assistant cannot do something (it may be wrong: Sheep can)
CANT = (r"\bi (?:can(?:no|')t|cannot|am not able to|'m not able to|am unable to|'m unable to|do not|don't)"
        r"(?: \w+){0,3}? (?:generate|create|make|draw|produce|render|record|speak|see|view|open|edit|send|"
        r"have the (?:capability|ability|tools?|means))|\bas an? (?:ai|text[- ]based|language) model\b|"
        r"\bi(?:'m| am) (?:just |only )?an? (?:text|language)[- ](?:based )?(?:model|ai)\b")


CLAIM = (r"\b(i('ve| have)? (added|created|made|updated|inserted|changed|saved|written|wrote|put)|"
         r"i('ll| will) (add|create|make|update|insert|put|change)|here'?s the (updated|new) (file|deck|document|"
         r"slides?|sheet|workbook|presentation)|has been (added|created|updated|saved|inserted)|"
         r"(added|created) (a|the) (new )?(slide|deck|document|sheet|section|row))\b|"
         # a reply that opens with the deed, no "I": "Added a bar chart to budget.ssheet.",
         # "Updated budget.ssheet: sorted by Cost" (the 3B, 2026-10-12 simulation; no tool ran)
         r"^\W*(added|created|made|updated|inserted|changed|saved|wrote|sorted|exported|renamed|moved|copied|"
         r"deleted|removed|formatted|replaced|rendered|charted|drafted|given|gave|set|applied|switched|turned)\b|"
         # "Your save-water.pptx is ready." with no export run (2026-10-13 simulation)
         r"\byour [\w./\\-]+\.(pptx|docx|xlsx|pdf|mp4|csv|md|sdoc|ssheet|sdeck|sstudio|eml|ics|txt|wav|png) "
         r"(is|are|has been) (ready|done|saved|created|exported|made)\b|"
         r"\b(has|have) been (sorted|exported|renamed|moved|copied|deleted|removed|formatted|replaced|rendered|charted)\b")


LONG_WAIT_S = 10


LONG_NOTES = {
    "make_image": "Making the picture", "edit_image": "Editing the picture", "studio_render": "Rendering the video", "transcribe_audio": "Writing down what is said",
    "speak_text": "Recording the voice", "web_fetch": "Reading the page", "web_search": "Searching the web",
    "export_file": "Exporting the file", "import_file": "Importing the file", "herder_run": "Running the tool",
}


ARTIFACT_TOOLS = {
    "slides_create", "slides_add", "slides_set", "slides_delete", "slides_move", "slides_set_theme",
    "words_create", "words_add", "words_add_list", "words_add_table", "words_replace", "words_delete",
    "sheets_create", "sheets_write", "sheets_set_cell", "sheets_add_sheet",
    "studio_create", "studio_add_media", "studio_add_text", "studio_add_shape", "studio_set",
    "studio_template", "studio_render", "make_image", "make_video",
}


# a detected job may use its own tools: "give me a powerpoint about water" was refused before (2026-10-20)
JOB_TOOLS = {"slides": {"slides_create", "slides_add", "export_file"},
             "words": {"words_create", "words_add", "words_add_list", "words_add_table", "export_file"},
             "sheets": {"sheets_create", "sheets_write", "sheets_set_cell", "sheets_add_sheet", "export_file"}}


# media tools the app checks reach but never run: running them lends the GPU under orch.lock or downloads models
REHEARSED = {"make_image", "make_video", "studio_render", "speak_text", "studio_voice", "transcribe_audio", "studio_transcribe"}


def error_words(e, local=False):
    """A model call that failed, in words a person can act on (2026-10-21: the chat showed
    "<urlopen error [WinError 10061] ...>" with no network and "401: User not found" for a wrong key)."""
    s = str(e)
    m = s.lower()
    # the local model's own failure is never Nebius's (J2, 2026-10-22: a compatibility server starting up)
    if local or "127.0.0.1" in m or "localhost" in m:
        return "The model on this computer stopped answering; Sheep is starting it again. Send again in a moment."
    if any(x in m for x in ("401", "403", "user not found", "invalid api key", "unauthorized", "incorrect api key")):
        return ("Nebius did not accept the key. Paste the right key in Settings › Accounts, "
                "then send again.")
    if "429" in m or "rate limit" in m:
        return "Nebius is busy right now. Send again in a minute."
    if any(x in m for x in ("urlopen error", "10061", "11001", "getaddrinfo", "timed out", "connection", "unreachable")):
        return ("Sheep could not reach Nebius. Check the internet connection, then send again. In Sheep mode the "
                "model on this computer answers with no connection.")
    if any(f" {c}" in f" {m}" for c in ("500", "502", "503", "504")):
        return "Nebius had a problem answering. Send again in a moment."
    return s.splitlines()[0][:200] if s else "The model did not answer."


def plain_error(r):
    """A tool's failure in words a person can act on."""
    s = str(r)
    return ("the file is open in another program; close it and try again" if ("WinError 5" in s or "PermissionError" in s)
            else s.split("\n")[0][:200])


def artifact_action_allowed(name, text):
    if name not in ARTIFACT_TOOLS or shepherd.explicit_change(text):
        return True

    # Reformatting the preceding answer as a file is an explicit artifact
    # request, even though it does not use one of the usual create verbs.
    # Keep the requested destination tied to its matching tool family so this
    # cannot authorize an unrelated artifact the model happened to suggest.
    target_tools = {
        "words": {"words_create", "words_add", "words_add_list", "words_add_table"},
        "sheets": {"sheets_create", "sheets_write", "sheets_set_cell", "sheets_add_sheet"},
        "slides": {"slides_create", "slides_add"},
        "video": {"make_video", "studio_render"},
    }
    destination = re.search(
        r"\bput\s+(?:that|this|it|them|those|these)\s+(?:in|into)\s+"
        r"(?:a|the|my)?\s*(document|doc|file|report|notes?|spreadsheet|workbook|sheet|"
        r"deck|slides?|presentation|slideshow|ppt|video|montage)\b",
        (text or "").lower(),
    )
    if not destination:
        return False
    noun = destination.group(1)
    family = (
        "words" if noun in {"document", "doc", "file", "report", "note", "notes"}
        else "sheets" if noun in {"spreadsheet", "workbook", "sheet"}
        else "slides" if noun in {"deck", "slide", "slides", "presentation", "slideshow", "ppt"}
        else "video"
    )
    return name in target_tools[family]


def long_note(name, args):
    """What a long step is doing, plus measured comparison data when available."""
    from ..tools import media
    if name == "make_video":
        return {"what": media.wait_words(args), "estimate_seconds": media.video_estimate_seconds(args)}
    if name == "make_image":
        return {"what": f"Making the picture on {media.PICTURES.words or 'this machine'}"}
    return {"what": LONG_NOTES.get(name, "Still working on " + name.replace("_", " "))}
