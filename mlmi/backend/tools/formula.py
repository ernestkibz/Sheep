"""Spreadsheet formulas, evaluated here so values show without Excel.

A saved .xlsx only carries the values Excel itself computed last time, and a
.ssheet carries none. So Sheep Sheets evaluates formulas: arithmetic, comparisons, & joins, cell references, ranges,
cross-sheet references, and the functions people actually use. The formula is
still saved as the formula, so Excel recomputes it the same way.

Anything outside this set evaluates to "#NAME?" -- visibly unknown, never a
silently wrong number.
"""
import math
import re

from datetime import date as _date, timedelta as _delta

from openpyxl.utils import column_index_from_string, get_column_letter

EPOCH = _date(1899, 12, 30)

TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<fn>[A-Za-z][A-Za-z0-9.]*)\s*\(
  | (?P<ref>(?:(?:'[^']+'|[A-Za-z_][\w.]*)!)?\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?)
  | (?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?%?|\.\d+%?)
  | (?P<bool>TRUE|FALSE)\b
  | (?P<op><=|>=|<>|[-+*/^&=<>(),:%])
""", re.X | re.I)


class FErr(Exception):
    def __init__(self, code):
        self.code = code


class _Bad:
    """An argument that failed: only an error if the function uses it (IF's other branch, IFERROR)."""
    def __init__(self, code):
        self.code = code


def _tokens(src):
    out, i = [], 0
    while i < len(src):
        m = TOKEN.match(src, i)
        if not m:
            raise FErr("#NAME?")
        kind = m.lastgroup
        if kind != "ws":
            out.append((kind, m.group(kind) if kind != "fn" else m.group("fn")))
        i = m.end()
    return out


def _num(v):
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?$", str(v))
    if m:                              # a date counts as Excel's day number: end - start works
        return float((_date(*map(int, m.groups())) - EPOCH).days)
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        raise FErr("#VALUE!")


def _text_as(v, fmt):
    """TEXT(): the number and date formats workbooks really use."""
    f = str(fmt)
    if re.search(r"[ymd]", f, re.I) and not re.search(r"[0#%]", f):
        d = EPOCH + _delta(days=int(_num(v)))
        parts = {"yyyy": f"{d.year}", "yy": f"{d.year % 100:02d}", "mmmm": d.strftime("%B"), "mmm": d.strftime("%b"),
                 "mm": f"{d.month:02d}", "m": str(d.month), "dddd": d.strftime("%A"), "ddd": d.strftime("%a"),
                 "dd": f"{d.day:02d}", "d": str(d.day)}
        return re.sub(r"y{4}|y{2}|m{1,4}|d{1,4}", lambda x: parts[x.group().lower()], f, flags=re.I)
    if not re.search(r"[0#]", f):
        return _text(v)
    n, dec = _num(v), len(f.split(".")[1].rstrip("%")) if "." in f else 0
    if "%" in f:
        return f"{n * 100:.{dec}f}%"
    return f"{n:,.{dec}f}" if "," in f else f"{n:.{dec}f}"


def _index(a):
    rng = a[0] if isinstance(a[0], list) else [[a[0]]]
    i, j = int(_num(a[1])), int(_num(a[2])) if len(a) > 2 else None
    if i < 1 and j is None or (j is not None and j < 1):
        raise FErr("#REF!")
    if j is None:
        j, i = (1, i) if len(rng) > 1 or len(rng[0]) == 1 else (i, 1)
    return rng[i - 1][j - 1]


def _match(a):
    flat, x = _flat([a[1]]), a[0]
    kind = int(_num(a[2])) if len(a) > 2 else 1
    low = str(x).lower() if isinstance(x, str) else x
    for i, v in enumerate(flat, 1):
        if kind == 0 and (str(v).lower() if isinstance(v, str) else v) == low:
            return float(i)
    if kind != 0:
        hits = [i for i, v in enumerate(flat, 1) if isinstance(v, (int, float)) and v <= _num(x)]
        if hits:
            return float(hits[-1])
    raise FErr("#N/A")


def _flat(args):
    for a in args:
        if isinstance(a, list):
            for row in a:
                for v in (row if isinstance(row, list) else [row]):
                    yield v
        else:
            yield a


def _nums(args, strict=False):
    out = []
    for a in args:
        if isinstance(a, list):                 # ranges skip text and blanks, as Excel does
            out.extend(float(v) for v in _flat([a]) if isinstance(v, (int, float)) and not isinstance(v, bool))
        else:
            out.append(_num(a))
    return out


def _crit(c):
    m = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", str(c))
    op, val = m.group(1) or "=", m.group(2)
    try:
        val = float(val)
    except ValueError:
        pass

    def test(v):
        if isinstance(val, float):
            try:
                x = _num(v)
            except FErr:
                return op == "<>"
        else:
            x, v2 = str(v if v is not None else "").lower(), val.lower()
            return {"=": x == v2, "<>": x != v2}.get(op, False)
        return {"=": x == val, "<>": x != val, "<": x < val, ">": x > val, "<=": x <= val, ">=": x >= val}[op]
    return test


def _round(x, d, mode):
    f = 10 ** int(d)
    if mode == "up":
        return math.ceil(abs(x) * f) / f * (1 if x >= 0 else -1)
    if mode == "down":
        return math.floor(abs(x) * f) / f * (1 if x >= 0 else -1)
    return math.floor(abs(x) * f + 0.5) / f * (1 if x >= 0 else -1)


FUNCS = {
    "SUM": lambda a: sum(_nums(a)),
    "AVERAGE": lambda a: (lambda n: sum(n) / len(n) if n else (_ for _ in ()).throw(FErr("#DIV/0!")))(_nums(a)),
    "MIN": lambda a: min(_nums(a) or [0]),
    "MAX": lambda a: max(_nums(a) or [0]),
    "COUNT": lambda a: float(len(_nums(a))),
    "COUNTA": lambda a: float(sum(1 for v in _flat(a) if v not in (None, ""))),
    "ABS": lambda a: abs(_num(a[0])),
    "ROUND": lambda a: _round(_num(a[0]), _num(a[1]) if len(a) > 1 else 0, "half"),
    "ROUNDUP": lambda a: _round(_num(a[0]), _num(a[1]) if len(a) > 1 else 0, "up"),
    "ROUNDDOWN": lambda a: _round(_num(a[0]), _num(a[1]) if len(a) > 1 else 0, "down"),
    "INT": lambda a: float(math.floor(_num(a[0]))),
    "SQRT": lambda a: math.sqrt(_num(a[0])),
    "POWER": lambda a: _num(a[0]) ** _num(a[1]),
    "MOD": lambda a: _num(a[0]) % _num(a[1]),
    "IF": lambda a: a[1] if _truth(a[0]) else (a[2] if len(a) > 2 else False),
    "AND": lambda a: all(_truth(v) for v in _flat(a)),
    "OR": lambda a: any(_truth(v) for v in _flat(a)),
    "NOT": lambda a: not _truth(a[0]),
    "CONCATENATE": lambda a: "".join(_text(v) for v in a),
    "CONCAT": lambda a: "".join(_text(v) for v in _flat(a)),
    "LEN": lambda a: float(len(_text(a[0]))),
    "UPPER": lambda a: _text(a[0]).upper(),
    "LOWER": lambda a: _text(a[0]).lower(),
    "TRIM": lambda a: " ".join(_text(a[0]).split()),
    "LEFT": lambda a: _text(a[0])[:int(_num(a[1])) if len(a) > 1 else 1],
    "RIGHT": lambda a: _text(a[0])[-int(_num(a[1])) if len(a) > 1 else -1:],
    "SUMIF": lambda a: sum(_num(v) for v, s in zip(_flat([a[2] if len(a) > 2 else a[0]]), _flat([a[0]]))
                           if _crit(a[1])(s) and isinstance(v, (int, float))),
    "COUNTIF": lambda a: float(sum(1 for v in _flat([a[0]]) if _crit(a[1])(v))),
    "AVERAGEIF": lambda a: (lambda xs: sum(xs) / len(xs) if xs else (_ for _ in ()).throw(FErr("#DIV/0!")))(
        [_num(v) for v, s in zip(_flat([a[2] if len(a) > 2 else a[0]]), _flat([a[0]]))
         if _crit(a[1])(s) and isinstance(v, (int, float))]),
    "MID": lambda a: _text(a[0])[int(_num(a[1])) - 1:int(_num(a[1])) - 1 + int(_num(a[2]))],
    "REPT": lambda a: _text(a[0]) * max(0, int(_num(a[1]))),
    "TEXT": lambda a: _text_as(a[0], a[1]),
    "INDEX": _index,
    "MATCH": _match,
    "IFERROR": lambda a: a[1] if isinstance(a[0], _Bad) or (isinstance(a[0], str) and a[0].startswith("#")) else a[0],
    "PI": lambda a: math.pi,
    "TODAY": lambda a: __import__("datetime").date.today().isoformat(),
}


def _truth(v):
    if isinstance(v, str):
        if v.upper() in ("TRUE", "FALSE"):
            return v.upper() == "TRUE"
        raise FErr("#VALUE!")
    return bool(_num(v))


def _text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


class Evaluator:
    """Evaluates every formula of a book, following references, catching cycles.

    `book` has .sheetnames and .get(sheet, coord) -> raw value (formats.SheetView)."""

    def __init__(self, book):
        self.book = book
        self.cache = {}
        self.busy = set()

    def value(self, sheet, coord):
        key = (sheet, coord.upper().replace("$", ""))
        if key in self.cache:
            return self.cache[key]
        if key in self.busy:
            raise FErr("#REF!")          # a cycle
        if sheet not in self.book.sheetnames:
            return "#REF!"
        v = self.book.get(sheet, key[1])
        if isinstance(v, str) and v.startswith("="):
            self.busy.add(key)
            try:
                v = self.formula(sheet, v[1:])
            except FErr as e:
                v = e.code
            except (ZeroDivisionError,):
                v = "#DIV/0!"
            except (ValueError, OverflowError, TypeError, IndexError):
                v = "#VALUE!"
            finally:
                self.busy.discard(key)
        self.cache[key] = v
        return v

    def formula(self, sheet, src):
        # a reference to another formula cell evaluates it mid-parse: keep ours
        saved = (getattr(self, "toks", None), getattr(self, "i", 0), getattr(self, "sheet", None))
        try:
            self.toks = _tokens(src)
            self.i = 0
            self.sheet = sheet
            v = self._cmp()
            if self.i != len(self.toks):
                raise FErr("#NAME?")
            return v[0][0] if isinstance(v, list) and len(v) == 1 and len(v[0]) == 1 else v
        finally:
            self.toks, self.i, self.sheet = saved

    # grammar: cmp -> cat ((=|<>|<|>|<=|>=) cat)*; cat -> add (& add)*; add -> mul ((+|-) mul)*
    #          mul -> pow ((*|/) pow)*; pow -> unary (^ unary)*; unary -> (-|+)* post; post -> atom %?
    def _peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def _eat(self, val=None):
        t = self._peek()
        if val is not None and t[1] != val:
            raise FErr("#NAME?")
        self.i += 1
        return t

    def _scalar(self, v):
        if isinstance(v, list):
            if len(v) == 1 and len(v[0]) == 1:
                return v[0][0]
            raise FErr("#VALUE!")
        return v

    def _cmp(self):
        a = self._cat()
        while self._peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self._eat()[1]
            b = self._cat()
            x, y = self._scalar(a), self._scalar(b)
            if isinstance(x, str) or isinstance(y, str):
                x, y = _text(x).lower(), _text(y).lower()
            else:
                x, y = _num(x), _num(y)
            a = {"=": x == y, "<>": x != y, "<": x < y, ">": x > y, "<=": x <= y, ">=": x >= y}[op]
        return a

    def _cat(self):
        a = self._add()
        while self._peek()[1] == "&":
            self._eat()
            a = _text(self._scalar(a)) + _text(self._scalar(self._add()))
        return a

    def _add(self):
        a = self._mul()
        while self._peek()[1] in ("+", "-"):
            op = self._eat()[1]
            b = self._mul()
            a = _num(self._scalar(a)) + _num(self._scalar(b)) if op == "+" else _num(self._scalar(a)) - _num(self._scalar(b))
        return a

    def _mul(self):
        a = self._pow()
        while self._peek()[1] in ("*", "/"):
            op = self._eat()[1]
            b = _num(self._scalar(self._pow()))
            if op == "/" and b == 0:
                raise FErr("#DIV/0!")
            a = _num(self._scalar(a)) * b if op == "*" else _num(self._scalar(a)) / b
        return a

    def _pow(self):
        a = self._unary()
        while self._peek()[1] == "^":
            self._eat()
            a = _num(self._scalar(a)) ** _num(self._scalar(self._unary()))
        return a

    def _unary(self):
        if self._peek()[1] in ("-", "+"):
            op = self._eat()[1]
            v = _num(self._scalar(self._unary()))
            return -v if op == "-" else v
        v = self._atom()
        if self._peek()[1] == "%":
            self._eat()
            v = _num(self._scalar(v)) / 100.0
        return v

    def _atom(self):
        kind, val = self._eat()
        if kind == "num":
            return float(val[:-1]) / 100.0 if val.endswith("%") else float(val)
        if kind == "str":
            return val[1:-1].replace('""', '"')
        if kind == "bool":
            return val.upper() == "TRUE"
        if kind == "ref":
            return self._ref(val)
        if kind == "fn":
            name = val.upper()             # the token already took its "("
            args = []
            if self._peek()[1] != ")":
                while True:
                    end, depth = self.i, 0            # where this argument ends, should it fail part way
                    while end < len(self.toks):
                        k, t = self.toks[end]
                        depth += k == "fn" or t == "("
                        if depth == 0 and t in (",", ")"):
                            break
                        depth -= t == ")" and depth > 0
                        end += 1
                    try:
                        args.append(self._cmp())
                    except (FErr, ZeroDivisionError, ValueError, TypeError, IndexError) as e:
                        args.append(_Bad(e.code if isinstance(e, FErr) else
                                         "#DIV/0!" if isinstance(e, ZeroDivisionError) else "#VALUE!"))
                        self.i = end
                    if self._peek()[1] == ",":
                        self._eat()
                        continue
                    break
            self._eat(")")
            fn = FUNCS.get(name)
            if not fn:
                raise FErr("#NAME?")
            bad = next((x for x in (args[:1] if name == "IF" else [] if name == "IFERROR" else args) if isinstance(x, _Bad)), None)
            if bad:
                raise FErr(bad.code)
            if name == "IF":
                args = [self._scalar(args[0])] + args[1:]
            v = fn(args)
            if isinstance(v, _Bad):
                raise FErr(v.code)
            return float(v) if isinstance(v, int) and not isinstance(v, bool) else v
        if val == "(":
            v = self._cmp()
            self._eat(")")
            return v
        raise FErr("#NAME?")

    def _ref(self, text):
        sheet = self.sheet
        if "!" in text:
            s, text = text.rsplit("!", 1)
            sheet = s.strip("'")
            if sheet not in self.book.sheetnames:
                raise FErr("#REF!")
        text = text.replace("$", "").upper()
        if ":" not in text:
            v = self.value(sheet, text)
            if isinstance(v, str) and v.startswith("#"):
                raise FErr(v)
            return v
        a, b = text.split(":")
        ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
        cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
        c0, c1 = sorted((column_index_from_string(ca), column_index_from_string(cb)))
        r0, r1 = sorted((int(ra), int(rb)))
        if (r1 - r0 + 1) * (c1 - c0 + 1) > 200_000:
            raise FErr("#VALUE!")
        return [[self.value(sheet, f"{get_column_letter(c)}{r}") for c in range(c0, c1 + 1)]
                for r in range(r0, r1 + 1)]
