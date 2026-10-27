"""Proof that the Sheep ecosystem does what it claims, file by file.

    python mlmi/tests/test_ecosystem.py

For each app: a model-style call makes the file, a hand edit (the same call the
editor makes when a person types) changes it, it is exported to the Microsoft
format, the export is opened by an independent library (python-docx, openpyxl,
python-pptx), re-imported, and compared. Then folders, moves and undo.
Video and Motion are not built and are not tested here.
"""
import os
import sys
import tempfile
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", ".."))
from mlmi.backend import tools as T  # noqa: E402
from mlmi.backend.workspace import Workspace  # noqa: E402

ok = 0


def check(what, cond, detail=""):
    global ok
    print(("  PASS  " if cond else "  FAIL  ") + what + (f"  ({detail})" if detail else ""))
    if not cond:
        raise SystemExit(1)
    ok += 1


def main():
    ws = Workspace(Path(tempfile.mkdtemp(prefix="mlmi-proof-")))
    run = lambda name, **a: T.run(ws, name, a, by="agent")       # noqa: E731
    hand = lambda name, **a: T.run(ws, name, a, by="user")       # noqa: E731

    print("Sheep Words")
    run("words_create", path="brief.sdoc", title="Brief")
    run("words_add", path="brief.sdoc", text="Goals", kind="heading", level=1)
    run("words_add_list", path="brief.sdoc", items=["ship the app", "win the race"])
    run("words_add_table", path="brief.sdoc", rows=[["Item", "Cost"], ["Laptop", "0"]])
    hand("words_put", path="brief.sdoc", index=2, block={"type": "bullet", "level": 0, "text": "ship the app **by hand**"})
    d = T.words.model("brief.sdoc", ws)
    check("made by the model, edited by hand", d["blocks"][2]["text"] == "ship the app **by hand**", d["blocks"][2]["text"])
    run("export_file", path="brief.sdoc")
    import docx
    doc = docx.Document(str(ws.path("brief.docx")))
    paras = [p.text for p in doc.paragraphs if p.text]
    check(".docx opens in python-docx", "Brief" in paras and any("by hand" in p for p in paras), f"{len(paras)} paragraphs")
    check(".docx keeps the bold", any(r.bold and "by hand" in r.text for p in doc.paragraphs for r in p.runs))
    check(".docx keeps the table", len(doc.tables) == 1 and doc.tables[0].cell(1, 0).text == "Laptop")
    ws.path("in").mkdir()
    Path(ws.path("in/brief.docx")).write_bytes(ws.path("brief.docx").read_bytes())
    run("import_file", path="in/brief.docx")
    back = T.words.model("in/brief.sdoc", ws)
    check(".docx imports back to Sheep Words", any("by hand" in (b.get("text") or "") for b in back["blocks"]))

    print("Sheep Sheets")
    run("sheets_create", path="budget.ssheet", headers=["Item", "Cost"], rows=[["Rent", 800], ["Food", 250]], total=True)
    hand("sheets_set_cell", path="budget.ssheet", cell="B3", value="300")
    m = T.sheets.model("budget.ssheet", ws)
    total = m["values"]["Sheet1"]["B4"]
    check("hand edit recalculates the total", float(total) == 1100, f"B4 = {total}")
    run("export_file", path="budget.ssheet")
    import openpyxl
    wb = openpyxl.load_workbook(str(ws.path("budget.xlsx")))
    sh = wb.active
    check(".xlsx opens in openpyxl with the formula", sh["B4"].value == "=SUM(B2:B3)" and sh["B3"].value == 300, f"B4 {sh['B4'].value}")
    Path(ws.path("in/budget.xlsx")).write_bytes(ws.path("budget.xlsx").read_bytes())
    run("import_file", path="in/budget.xlsx")
    m2 = T.sheets.model("in/budget.ssheet", ws)
    check(".xlsx imports back with a live formula", float(m2["values"]["Sheet1"]["B4"]) == 1100)

    print("Sheep Presentations")
    run("slides_create", path="pitch.sdeck", title="Sheep", subtitle="local AI")
    run("slides_add", path="pitch.sdeck", title="Speed", bullets=["faster", "  on a laptop"])
    hand("slides_put", path="pitch.sdeck", slide=2, data={"layout": "bullets", "title": "Speed (edited)", "bullets": ["faster", "  on a laptop"],
                                                            "notes": "say it slowly", "items": []})
    run("slides_set_theme", path="pitch.sdeck", theme="midnight")
    run("export_file", path="pitch.sdeck")
    import pptx
    pr = pptx.Presentation(str(ws.path("pitch.pptx")))
    texts = [sh.text_frame.text for s in pr.slides for sh in s.shapes if sh.has_text_frame]
    check(".pptx opens in python-pptx", len(pr.slides) == 2 and any("Speed (edited)" in t for t in texts), f"{len(pr.slides)} slides")
    check(".pptx keeps the speaker notes", pr.slides[1].notes_slide.notes_text_frame.text == "say it slowly")
    Path(ws.path("in/pitch.pptx")).write_bytes(ws.path("pitch.pptx").read_bytes())
    run("import_file", path="in/pitch.pptx")
    dk = T.slides.model("in/pitch.sdeck", ws)
    check(".pptx imports back to Sheep Presentations", dk["slides"][1]["title"] == "Speed (edited)")

    print("Files and folders")
    run("create_folder", path="reports")
    run("move_file", src="brief.docx", dst="reports/brief.docx")
    run("copy_file", src="budget.xlsx", dst="reports/budget copy.xlsx")
    names = [f["path"] for f in ws.tree()]
    check("folder, move and copy", "reports/brief.docx" in names and "reports/budget copy.xlsx" in names and "brief.docx" not in names)
    ws.undo()
    check("undo puts the copy back out", "reports/budget copy.xlsx" not in [f["path"] for f in ws.tree()])
    print(f"all {ok} checks passed  (workspace: {ws.root})")


if __name__ == "__main__":
    main()
