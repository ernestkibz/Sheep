"""Every tool, graded by the file it leaves behind.

Each check acts through the registry -- the same path the agent and the
editors use -- then reads the result back: Sheep files as their JSON, exported
Office files with the library that reads Office files. The file is the grade.

    python mlmi/tests/test_tools.py
"""
import json
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
os.environ["MLMI_HOME"] = tempfile.mkdtemp(prefix="mlmi_home_")     # herders made here stay out of ~/.mlmi

from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

from mlmi.backend import tools as T  # noqa: E402
from mlmi.backend.workspace import ToolError, Workspace  # noqa: E402


def ws():
    d = tempfile.mkdtemp(prefix="mlmi_t_")
    return Workspace(d), d


def js(d, rel):
    with open(os.path.join(d, rel), encoding="utf-8") as f:
        return json.load(f)


def refused(fn):
    try:
        fn()
    except ToolError as e:
        return str(e)
    raise AssertionError("was allowed")


def test_paths_stay_inside():
    w, d = ws()
    for bad in ("../x.txt", "..\\..\\x.txt", "C:/Windows/x.txt", ".mlmi/journal.jsonl"):
        refused(lambda: T.run(w, "write_text_file", {"path": bad, "text": "x"}))
    shutil.rmtree(d)


def test_files_and_undo():
    w, d = ws()
    T.run(w, "write_text_file", {"path": "notes/a.txt", "text": "one"})
    T.run(w, "write_text_file", {"path": "notes/a.txt", "text": "two"})
    assert open(os.path.join(d, "notes/a.txt")).read() == "two"
    w.undo()
    assert open(os.path.join(d, "notes/a.txt")).read() == "one"
    w.redo()
    assert open(os.path.join(d, "notes/a.txt")).read() == "two"
    T.run(w, "delete_file", {"path": "notes/a.txt"})
    w.undo()
    assert open(os.path.join(d, "notes/a.txt")).read() == "two"
    T.run(w, "move_file", {"src": "notes/a.txt", "dst": "b.txt"})
    w.undo()
    assert os.path.exists(os.path.join(d, "notes/a.txt")) and not os.path.exists(os.path.join(d, "b.txt"))
    assert T.needs_confirm(w, "delete_file", {"path": "notes/a.txt"})
    assert T.needs_confirm(w, "write_text_file", {"path": "notes/a.txt", "text": "x"})
    assert not T.needs_confirm(w, "write_text_file", {"path": "new.txt", "text": "x"})
    shutil.rmtree(d)


def test_words():
    w, d = ws()
    T.run(w, "words_create", {"path": "report.docx", "title": "Quarterly Report"})    # asks for Word, gets Sheep
    assert os.path.exists(os.path.join(d, "report.sdoc"))
    T.run(w, "words_add", {"path": "report.sdoc", "text": "Summary", "kind": "heading", "level": 1})
    T.run(w, "words_add", {"path": "report.sdoc", "text": "Sales rose **12%** this *quarter*."})
    # the live-run failure: a list passed as a Python-looking string, a bool as text
    T.run(w, "words_add_list", {"path": "report.sdoc", "items": "['North', 'South']", "numbered": "false"})
    T.run(w, "words_add_table", {"path": "report.sdoc", "rows": [["Region", "Sales"], ["North", "120"], ["South", 95]]})
    T.run(w, "words_add", {"path": "report.sdoc", "text": "Intro", "kind": "heading", "level": "2", "after": "0"})
    out = T.run(w, "words_read", {"path": "report.sdoc"})
    assert "[1] heading 2: Intro" in out and "[4] bullet: North" in out and "[5] bullet: South" in out, out
    T.run(w, "words_replace", {"path": "report.sdoc", "find": "Sales", "replace": "Revenue"})
    T.run(w, "words_set_cell", {"path": "report.sdoc", "index": 6, "row": 1, "col": 1, "text": "130"})
    doc = js(d, "report.sdoc")
    assert doc["blocks"][6]["rows"][1][1] == "130" and doc["blocks"][6]["rows"][0][1] == "Revenue"
    T.run(w, "words_delete", {"path": "report.sdoc", "index": 1})
    assert "Intro" not in json.dumps(js(d, "report.sdoc"))
    w.undo()
    assert "Intro" in json.dumps(js(d, "report.sdoc"))
    # A Microsoft name means its Sheep copy -- the live-run failure: the model made
    # shopping.docx (-> .sdoc) and went on calling it shopping.docx
    rc = T.Ctx(w, "agent")
    assert "already exists" in T.run(w, "words_create", {"path": "report.docx"}, ctx=rc)    # no silent wipe
    out = T.run(w, "words_add", {"path": "report.docx", "text": "Added via the Word name"}, ctx=rc)
    assert "report.sdoc" in out and "Added via" in json.dumps(js(d, "report.sdoc"))
    assert rc.exports == {"report.sdoc": "report.docx"}           # exported when the job ends
    assert "words_create" in refused(lambda: T.run(w, "words_read", {"path": "nothing.docx"}))
    # an Office file with no Sheep copy yet is imported on first use
    T.run(w, "export_file", {"path": "report.sdoc", "to": "other.docx"})
    assert "imported other.docx" in T.run(w, "words_read", {"path": "other.docx"}, ctx=rc)
    assert os.path.exists(os.path.join(d, "other.sdoc"))
    # a file this job wrote can be replaced without asking; someone else's cannot
    assert T.needs_confirm(w, "export_file", {"path": "report.sdoc", "to": "other.docx"})
    rc.made.add("other.docx")
    assert not T.needs_confirm(w, "export_file", {"path": "report.sdoc", "to": "other.docx"}, rc)
    # import asks nothing about a source that does not exist; the tool then says so
    assert not T.needs_confirm(w, "import_file", {"path": "ghost.docx"})
    shutil.rmtree(d)


def test_words_round_trip():
    w, d = ws()
    T.run(w, "words_create", {"path": "r.sdoc", "title": "Plan"})
    T.run(w, "words_add", {"path": "r.sdoc", "text": "Goals", "kind": "heading"})
    T.run(w, "words_add", {"path": "r.sdoc", "text": "We **will** ship *soon*."})
    T.run(w, "words_add_list", {"path": "r.sdoc", "items": ["one", "two"], "numbered": True})
    T.run(w, "words_add_table", {"path": "r.sdoc", "rows": [["A", "B"], ["1", "2"]]})
    T.run(w, "export_file", {"path": "r.sdoc"})
    doc = Document(os.path.join(d, "r.docx"))
    texts = [p.text for p in doc.paragraphs]
    assert texts[:3] == ["Plan", "Goals", "We will ship soon."], texts
    par = doc.paragraphs[2]
    assert any(r.bold and r.text == "will" for r in par.runs) and any(r.italic and r.text == "soon" for r in par.runs)
    assert doc.tables[0].cell(1, 1).text == "2"
    T.run(w, "import_file", {"path": "r.docx", "to": "back.sdoc"})
    back = js(d, "back.sdoc")["blocks"]
    kinds = [b["type"] for b in back]
    assert kinds == ["heading", "heading", "paragraph", "numbered", "numbered", "table"], kinds
    assert back[2]["text"] == "We **will** ship *soon*." and back[5]["rows"][1] == ["1", "2"]
    T.run(w, "export_file", {"path": "r.sdoc", "format": "md"})
    md = open(os.path.join(d, "r.md"), encoding="utf-8").read()
    assert "# Goals" in md and "1. one" in md and "| A | B |" in md
    shutil.rmtree(d)


def test_sheets():
    w, d = ws()
    T.run(w, "sheets_create", {"path": "b.ssheet", "sheet": "Budget"})
    T.run(w, "sheets_write", {"path": "b.ssheet", "cell": "A1", "values": [
        ["Item", "Cost"], ["Rent", 1200], ["Food", "350"], ["Travel", 90.5],
        ["Total", "=SUM(B2:B4)"], ["Avg", "=ROUND(AVERAGE(B2:B4),1)"], ["Big", "=COUNTIF(B2:B4,\">100\")"],
        ["Label", "=A2&\" & \"&A3"], ["Check", "=IF(B5>1500,\"over\",\"ok\")"]]})
    out = T.run(w, "sheets_read", {"path": "b.ssheet"})
    assert "1640.5 [=SUM(B2:B4)]" in out, out
    assert "546.8" in out and "\t2 [=COUNTIF" in out and "Rent & Food" in out and "over" in out, out
    assert js(d, "b.ssheet")["sheets"][0]["cells"]["B3"] == 350
    T.run(w, "sheets_set_cell", {"path": "b.ssheet", "cell": "B3", "value": "400"})
    assert "1690.5 [=SUM" in T.run(w, "sheets_read", {"path": "b.ssheet"})
    T.run(w, "sheets_add_sheet", {"path": "b.ssheet", "name": "Q2"})
    T.run(w, "sheets_write", {"path": "b.ssheet", "sheet": "Q2", "cell": "A1", "values": [["=Budget!B5*2"]]})
    assert "3381" in T.run(w, "sheets_read", {"path": "b.ssheet", "sheet": "Q2"})
    T.run(w, "sheets_write", {"path": "b.ssheet", "sheet": "Q2", "cell": "C1", "values": [["=D1", "=C1"]]})
    assert "#REF!" in T.run(w, "sheets_read", {"path": "b.ssheet", "sheet": "Q2"})
    T.run(w, "sheets_format", {"path": "b.ssheet", "range": "A1:B1", "bold": True, "fill": "DDEBF7", "width": 14})
    T.run(w, "sheets_add_chart", {"path": "b.ssheet", "range": "A1:B4", "kind": "bar", "title": "Costs"})
    T.run(w, "export_file", {"path": "b.ssheet"})
    wb = load_workbook(os.path.join(d, "b.xlsx"))
    assert wb["Budget"]["B5"].value == "=SUM(B2:B4)" and wb["Budget"]["B3"].value == 400
    assert wb["Budget"]["A1"].font.b and wb["Budget"].column_dimensions["A"].width == 14
    assert len(wb["Budget"]._charts) == 1
    T.run(w, "import_file", {"path": "b.xlsx", "to": "back.ssheet"})
    back = js(d, "back.ssheet")["sheets"][0]
    assert back["cells"]["B5"] == "=SUM(B2:B4)" and back["fmt"]["A1"]["b"] is True
    T.run(w, "export_file", {"path": "b.ssheet", "format": "csv"})
    assert "Total,1690.5" in open(os.path.join(d, "b.csv"), encoding="utf-8").read()
    shutil.rmtree(d)


def test_slides():
    w, d = ws()
    T.run(w, "slides_create", {"path": "deck.sdeck", "title": "Launch Plan", "subtitle": "2026", "theme": "midnight"})
    T.run(w, "slides_add", {"path": "deck.sdeck", "title": "Goals", "bullets": ["Grow", "  in Kenya", "Hire"]})
    T.run(w, "slides_add", {"path": "deck.sdeck", "title": "Budget", "bullets": ["$10k"], "notes": "keep it short"})
    T.run(w, "slides_add", {"path": "deck.sdeck", "title": "Agenda", "bullets": ["one"], "after": 1})
    out = T.run(w, "slides_read", {"path": "deck.sdeck"})
    assert "slide 2: Agenda" in out and "slide 3: Goals" in out and "notes: keep it short" in out, out
    T.run(w, "export_file", {"path": "deck.sdeck"})
    prs = Presentation(os.path.join(d, "deck.pptx"))
    assert len(prs.slides) == 4 and abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    goals = prs.slides[2]
    body = [sh for sh in goals.placeholders if sh.placeholder_format.idx == 1][0]
    assert [p.level for p in body.text_frame.paragraphs] == [0, 1, 0]
    assert str(goals.background.fill.fore_color.rgb) == "0B1020"
    t = prs.slides[0].shapes.title
    assert t.left + t.width > prs.slide_width * 0.85
    T.run(w, "import_file", {"path": "deck.pptx", "to": "back.sdeck"})
    back = js(d, "back.sdeck")
    assert [s["title"] for s in back["slides"]] == ["Launch Plan", "Agenda", "Goals", "Budget"]
    assert back["slides"][2]["bullets"] == ["Grow", "  in Kenya", "Hire"] and back["theme"] == "midnight"
    assert back["slides"][3]["notes"] == "keep it short"
    T.run(w, "slides_move", {"path": "deck.sdeck", "slide": 4, "to": 1})
    T.run(w, "slides_delete", {"path": "deck.sdeck", "slide": 2})
    assert [s["title"] for s in js(d, "deck.sdeck")["slides"]] == ["Budget", "Agenda", "Goals"]
    shutil.rmtree(d)


def test_herders():
    w, d = ws()
    # the built-ins run on real documents
    T.run(w, "words_create", {"path": "r.sdoc", "title": "Doc"})
    for t, lv in (("Intro", 1), ("Scope", 2), ("Detail", 2), ("Plan", 1)):
        T.run(w, "words_add", {"path": "r.sdoc", "text": t, "kind": "heading", "level": lv})
    T.run(w, "h_number_headings", {"path": "r.sdoc"})
    T.run(w, "herder_run", {"name": "number_headings", "path": "r.sdoc"})           # twice: no double numbers
    heads = [b["text"] for b in js(d, "r.sdoc")["blocks"][1:]]
    assert heads == ["1. Intro", "1.1 Scope", "1.2 Detail", "2. Plan"], heads
    w.undo()
    T.run(w, "sheets_create", {"path": "s.ssheet"})
    T.run(w, "sheets_write", {"path": "s.ssheet", "cell": "A1", "values": [["Item", "Q1", "Q2"], ["a", 1, 2], ["b", 3, 4]]})
    T.run(w, "h_add_totals", {"path": "s.ssheet", "range": "A1:C3"})
    assert "4 [=SUM(B2:B3)]" in T.run(w, "sheets_read", {"path": "s.ssheet"})
    T.run(w, "slides_create", {"path": "p.sdeck", "title": "Deck"})
    T.run(w, "slides_add", {"path": "p.sdeck", "title": "Why"})
    T.run(w, "slides_add", {"path": "p.sdeck", "title": "How"})
    T.run(w, "h_agenda_slide", {"path": "p.sdeck"})
    sl = js(d, "p.sdeck")["slides"]
    assert sl[1]["title"] == "Agenda" and sl[1]["bullets"] == ["Why", "How"]
    # the model writes one of its own, and it becomes a tool
    code = ("def herd(doc, args):\n"
            "    n = 0\n"
            "    for b in doc['blocks']:\n"
            "        if b.get('type') == 'paragraph':\n"
            "            b['text'] = b['text'].upper(); n += 1\n"
            "    print('seen', len(doc['blocks']))\n"
            "    return f'shouted {n} paragraphs'\n")
    T.run(w, "words_add", {"path": "r.sdoc", "text": "quiet words"})
    msg = T.run(w, "herder_make", {"name": "shout", "app": "words", "description": "uppercase paragraphs", "code": code})
    assert "h_shout" in msg and "h_shout" in T.TOOLS
    out = T.run(w, "h_shout", {"path": "r.sdoc"})
    assert "shouted 1 paragraphs" in out and "seen" in out, out
    assert js(d, "r.sdoc")["blocks"][-1]["text"] == "QUIET WORDS"
    assert "shout" in T.run(w, "herder_list", {"app": "words"})
    # what a herder may not do
    for bad in ("import os\ndef herd(doc, args):\n    os.remove('x')\n",
                "def herd(doc, args):\n    open('x','w').write('1')\n",
                "def herd(doc, args):\n    return ().__class__.__base__\n",
                "def herd(doc, args):\n    eval('1')\n",
                "def nothing():\n    pass\n"):
        refused(lambda: T.run(w, "herder_make", {"name": "evil", "app": "words", "description": "x", "code": bad}))
    # a runtime failure names its line; an endless loop is stopped
    T.run(w, "herder_make", {"name": "broken", "app": "words", "description": "x",
                             "code": "def herd(doc, args):\n    x = 1\n    return doc['nope']\n"})
    assert "line 3" in refused(lambda: T.run(w, "h_broken", {"path": "r.sdoc"}))
    import mlmi.backend.tools.herders as H
    H.TIMEOUT = 3
    T.run(w, "herder_make", {"name": "spin", "app": "words", "description": "x",
                             "code": "def herd(doc, args):\n    while True:\n        pass\n"})
    assert "stopped" in refused(lambda: T.run(w, "h_spin", {"path": "r.sdoc"}))
    T.run(w, "herder_delete", {"name": "spin"})
    assert "h_spin" not in T.TOOLS
    shutil.rmtree(d)


def test_registry():
    names = T.offered(1)
    assert "words_read" in names and "words_add" not in names and "herder_help" in names
    assert "delete_file" not in T.offered(2) and "delete_file" in T.offered(3)
    assert "words_put" not in T.offered(3)
    for s in T.schemas(T.offered(3)):
        f = s["function"]
        assert f["name"] and f["description"] and f["parameters"]["type"] == "object"
    assert "no tool named" in refused(lambda: T.run(Workspace(tempfile.mkdtemp()), "no_such_tool", {}))


if __name__ == "__main__":
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"  PASS  {name}")
            except Exception as e:   # noqa: BLE001
                fails += 1
                import traceback
                print(f"  FAIL  {name}: {e}")
                traceback.print_exc()
    print(f"{'all passed' if not fails else str(fails) + ' failed'}")
    sys.exit(1 if fails else 0)
