"""What the 2026-10-14 live run with the 3B found, kept fixed: pictures, voice, photo edits
and Office decks asked in plain words; and the Unigram tokenizer (Helium, XGLM, Persimmon)."""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tooling"))

from mlmi.backend import agent, shepherd, tools as T  # noqa: E402
from mlmi.backend.workspace import Workspace  # noqa: E402


class CreateAsks(unittest.TestCase):
    def test_pictures_voice_and_edits_are_asks_to_make(self):
        for t in ["Draw me a picture of a dog playing in the snow.", "a picture of a cat please",
                  "Read this sentence out loud in a voice: Welcome to the bakery.",
                  "Edit logo.png to make it black and white."]:
            self.assertTrue(shepherd.explicit_change(t), t)
        for t in ["what is a dog", "How do I draw a dog?"]:
            self.assertFalse(shepherd.explicit_change(t), t)

    def test_words_after_a_colon_are_spoken(self):
        self.assertEqual(agent.words_to_speak("Read this sentence out loud in a voice: Welcome to the bakery."),
                         "Welcome to the bakery.")
        self.assertEqual(agent.words_to_speak('Say "hello there" out loud'), "hello there")
        self.assertTrue(agent.explicit_speech_output("Read this out loud: hi"))

    def test_office_deck_slides_go_into_the_deck(self):
        self.assertEqual(shepherd.native("save-water.pptx", "slides"), "save-water.sdeck")
        self.assertEqual(shepherd.native("notes.docx", "words"), "notes.sdoc")
        self.assertEqual(shepherd.native("x.sdeck", "slides"), "x.sdeck")


class PhotoEdits(unittest.TestCase):
    def test_ops_read_from_plain_words(self):
        self.assertEqual(agent.image_ops_in("Edit logo.png to make it black and white."), "black and white")
        self.assertEqual(agent.image_ops_in("turn logo.png upside down"), "rotate 180")
        self.assertEqual(agent.image_ops_in("put logo.png on slide 2"), "")

    def test_edit_image_keeps_the_original(self):
        from PIL import Image
        ws = Workspace(Path(tempfile.mkdtemp()))
        Image.new("RGB", (40, 20), (200, 30, 30)).save(ws.root / "logo.png")
        out = T.run(ws, "edit_image", {"path": "logo.png", "ops": "black and white, rotate 90"}, by="test")
        self.assertIn("logo-edited.png", out)
        im = Image.open(ws.root / "logo-edited.png")
        self.assertEqual(im.size, (20, 40))
        r, g, b = im.convert("RGB").getpixel((5, 5))
        self.assertTrue(r == g == b)
        self.assertEqual(Image.open(ws.root / "logo.png").getpixel((0, 0)), (200, 30, 30))

    def test_photo_edit_routes_away_from_slides(self):
        names = agent.route_tools(set(T.TOOLS), "", "Edit logo.png to make it black and white.")
        self.assertIn("edit_image", names)
        self.assertNotIn("slides_add_image", names)


class Unigram(unittest.TestCase):
    def info(self, flags, kinds=None):
        import numpy as np
        from sml_unigram import NORMAL
        toks = ["<unk>", "▁", "▁the", "▁cat", "c", "a", "t", "h", "e", "▁t"]
        scores = np.asarray([0, -2, -1, -1.5, -5, -5, -5, -5, -5, -3], dtype=np.float32)
        k = np.full(len(toks), NORMAL, dtype=np.uint8) if kinds is None else kinds
        return toks, {"unk": 0, "flags": flags, "scores": scores, "kinds": k, "unk_score": -15.0,
                      "max_score": -1.0, "norm": []}

    def test_best_path(self):
        from sml_unigram import PREFIX_ALWAYS, SPLIT_WORDS, Unigram as U
        toks, info = self.info(PREFIX_ALWAYS | SPLIT_WORDS)
        self.assertEqual(U(toks, info).encode("the cat"), [2, 3])
        self.assertEqual(U(toks, info).encode("the dog")[:1], [2])       # unknowns fused into <unk>

    def test_trailer_round_trip(self):
        import struct
        from sml_format import build_vocab_blob
        from sml_unigram import PREFIX_ALWAYS
        toks, info = self.info(PREFIX_ALWAYS)
        blob = build_vocab_blob(toks, [], 0, 0, info)
        n, m = struct.unpack_from("<2I", blob, 0)
        end = 16 + 4 * (n + 1) + sum(len(t.encode()) for t in toks) + 8 * m
        self.assertEqual(blob[end:end + 4], b"UNIG")
        self.assertEqual(struct.unpack_from("<I", blob, end + 12)[0], PREFIX_ALWAYS)


class ConfigNames(unittest.TestCase):
    def test_fairseq_head_count_is_read(self):
        from sml_hf import _CFG_ALIASES
        self.assertIn("attention_heads", _CFG_ALIASES["num_attention_heads"])
        self.assertIn("ffn_dim", _CFG_ALIASES["intermediate_size"])


if __name__ == "__main__":
    unittest.main()


class TitleSize(unittest.TestCase):
    def test_bigger_title_reaches_the_slide_and_pptx(self):
        ws = Workspace(Path(tempfile.mkdtemp()))
        T.run(ws, "slides_create", {"path": "d.sdeck", "title": "Save Water"}, by="test")
        out = T.run(ws, "slides_set", {"path": "d.sdeck", "slide": 1, "title_size": "bigger"}, by="test")
        self.assertIn("1.25x", out)
        T.run(ws, "slides_set", {"path": "d.sdeck", "slide": 0, "title_size": "bigger"}, by="test")
        import json
        d = json.loads((ws.root / "d.sdeck").read_text(encoding="utf-8"))
        self.assertAlmostEqual(d["slides"][0]["title_scale"], 1.562, places=2)
        T.run(ws, "export_file", {"path": "d.sdeck", "format": "pptx"}, by="test")
        from pptx import Presentation
        p = Presentation(str(ws.root / "d.pptx"))
        sizes = [r.font.size.pt for r in p.slides[0].shapes.title.text_frame.paragraphs[0].runs if r.font.size]
        self.assertEqual(sizes, [69.0])

    def test_one_line_title_command(self):
        from mlmi.backend import direct
        self.assertEqual(direct.parse("Make the title bigger in save-water.sdeck")[:2],
                         ("slides_set", {"path": "save-water.sdeck", "slide": 1, "title_size": "bigger"}))


class SlideWithoutTitle(unittest.TestCase):
    """A planned slides_add with no title stopped a research deck after its title slide (J4, 2026-10-21)."""
    def test_first_bullet_becomes_the_title_or_it_is_numbered(self):
        ws = Workspace(Path(tempfile.mkdtemp()))
        T.run(ws, "slides_create", {"path": "d.sdeck", "title": "Solar"}, by="test")
        T.run(ws, "slides_add", {"path": "d.sdeck", "bullets": ["Growth", "Panels on roofs"]}, by="test")
        T.run(ws, "slides_add", {"path": "d.sdeck"}, by="test")
        titles = [t for t, _ in shepherd.deck_slides(ws, "d.sdeck")]
        self.assertEqual(titles[1:], ["growth", "slide 3"])


class FileTitles(unittest.TestCase):
    def test_request_format_words_are_not_the_title(self):
        self.assertEqual(shepherd.clean_title("Powerpoint about the History of Coffee"), "The History of Coffee")
        self.assertEqual(shepherd.clean_title("A presentation on solar power"), "Solar power")
        self.assertEqual(shepherd.clean_title("Coffee Report"), "Coffee Report")
