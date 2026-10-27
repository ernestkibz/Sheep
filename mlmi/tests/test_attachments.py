"""Attachments: saved under attachments/ without overwriting, then made readable
(Word imported, PDF text, a recording transcribed, a picture kept).

    python -m mlmi.tests.test_attachments
"""
import io
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

if __name__ != "__main__":
    raise unittest.SkipTest("CLI smoke test; run python -m mlmi.tests.test_attachments")

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from mlmi.backend import attachments as A  # noqa: E402
from mlmi.backend import tools as T  # noqa: E402
from mlmi.backend.workspace import Workspace  # noqa: E402

fails = []


def check(name, ok, got=""):
    print(("PASS  " if ok else "FAIL  ") + name + ("" if ok else f"   got {got!r}"))
    if not ok:
        fails.append(name)


root = Path(tempfile.mkdtemp(prefix="mlmi-attach-"))
ws = Workspace(root)
try:
    from docx import Document
    d = Document()
    d.add_heading("Quarterly plan", 0)
    d.add_paragraph("Sales rose by twelve percent in the third quarter.")
    buf = io.BytesIO()
    d.save(buf)
    p1 = A.save(ws, "plan.docx", buf.getvalue())
    p2 = A.save(ws, "plan.docx", buf.getvalue())
    check("saved under attachments/", p1 == "attachments/plan.docx", p1)
    check("a second file of that name is not overwritten", p2 == "attachments/plan-2.docx", p2)

    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DecodedStreamObject, DictionaryObject
    w = PdfWriter()
    page = w.add_blank_page(width=300, height=200)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 100 Td (The meeting is on Friday at ten.) Tj ET")
    page[NameObject("/Contents")] = w._add_object(stream)
    buf = io.BytesIO()
    w.write(buf)
    p3 = A.save(ws, "notice.pdf", buf.getvalue())
    p4 = A.save(ws, "photo.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64)

    ctx = T.Ctx(ws, "agent")
    got = dict(A.intake(ws, [p1, p3, p4], ctx))
    check("Word is imported into Sheep", "attachments/plan.sdoc" in got and ws.path("attachments/plan.sdoc").exists(), got)
    txt = ws.path("attachments/notice.txt")
    check("PDF text extracted", txt.exists() and "Friday" in txt.read_text(encoding="utf-8"),
          txt.read_text(encoding="utf-8") if txt.exists() else "missing")
    check("a picture is kept as it is", p4 in got, got)

    if "--audio" in sys.argv:
        T.run(ws, "speak_text", {"text": "Please send the report by Monday.", "path": "voice.wav"}, by="test", ctx=ctx)
        shutil.move(str(ws.path("voice.wav")), str(ws.path("attachments/voice.wav")))
        got = dict(A.intake(ws, ["attachments/voice.wav"], ctx))
        t = ws.path("attachments/voice.txt")
        check("a recording is transcribed", t.exists() and "monday" in t.read_text(encoding="utf-8").lower(),
              t.read_text(encoding="utf-8") if t.exists() else got)
finally:
    shutil.rmtree(root, ignore_errors=True)

print(f"\n{'all checks passed' if not fails else str(len(fails)) + ' failed: ' + ', '.join(fails)}")
sys.exit(1 if fails else 0)
